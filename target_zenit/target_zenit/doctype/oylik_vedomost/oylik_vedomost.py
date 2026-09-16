# Copyright (c) 2026, abdulloh and contributors
# For license information, please see license.txt

"""Oylik vedomost — har oylik ish haqi jadvali (buxgalter Excel faylidan yuklanadi).

Excel har oy o'zgaradi, shuning uchun ustunlar joyiga emas SARLAVHA NOMIGA qarab
topiladi, varaq esa avtomatik aniqlanadi (Ф.И.О. va Оклад ustunlari bor varaq).
Xodim bazadagi Employee bilan mos kelmasa ham qator yo'qolmaydi — F.I.Sh
vedomostdagi holicha saqlanadi.
"""

import re

import frappe
from frappe import _
from frappe.model.document import Document
from frappe.utils import cint, flt

# Excel sarlavhasi -> maydon. Kalitlar kichik harfda, bo'shliqsiz solishtiriladi.
COLUMN_MAP = {
	"ф.и.о.": "fio", "фио": "fio", "f.i.sh": "fio", "fish": "fio",
	"отдел": "otdel", "bo'lim": "otdel", "bolim": "otdel",
	"дожность": "doljnost", "должность": "doljnost", "lavozim": "doljnost",
	"тип": "tip", "tip": "tip",
	"окладфин": "oklad", "оклад": "oklad", "oklad": "oklad",
	"день": "kun", "kun": "kun",
	"режимдня": "rejim", "rejim": "rejim",
	"начисление": "nachisleniya", "nachisleniya": "nachisleniya",
	"бонус": "bonus", "bonus": "bonus",
	"квартира": "kvartira", "kvartira": "kvartira",
	"налог12%": "nalog", "налог": "nalog", "soliq": "nalog",
	"удержание(итого)": "ushlanma", "удержание": "ushlanma", "ushlanma": "ushlanma",
	"банк": "bank", "bank": "bank",
	"пластик": "plastik", "plastik": "plastik",
	"наличка": "nalichka", "naqd": "nalichka",
	"остатка": "qoldiq", "остаток": "qoldiq", "qoldiq": "qoldiq",
}

NUMERIC_FIELDS = ("oklad", "kun", "rejim", "nachisleniya", "bonus", "kvartira",
                  "nalog", "ushlanma", "bank", "plastik", "nalichka", "qoldiq")


def _norm_header(v):
	return re.sub(r"\s+", "", str(v or "").strip().lower())


def _norm_name(v):
	"""Ismni solishtirish uchun normallashtirish (apostrof/registr/ortiqcha bo'shliq)."""
	s = str(v or "").lower().replace("`", "").replace("'", "").replace("’", "")
	return " ".join(re.sub(r"[^a-z0-9Ѐ-ӿ ]", " ", s).split())


class OylikVedomost(Document):
	def validate(self):
		self.set_totals()
		self.link_employees()

	def set_totals(self):
		self.jami_oklad = sum(flt(r.oklad) for r in self.qatorlar)
		self.jami_nachisleniya = sum(flt(r.nachisleniya) for r in self.qatorlar)
		self.jami_qoldiq = sum(flt(r.qoldiq) for r in self.qatorlar)
		self.xodimlar_soni = len(self.qatorlar)

	def link_employees(self):
		"""Bo'sh qolgan employee maydonini ism bo'yicha AYNAN moslikda to'ldirish.
		Taxminiy moslashtirilmaydi — noto'g'ri bog'lash raqamlarni buzadi."""
		empty = [r for r in self.qatorlar if not r.employee and r.fio]
		if not empty:
			return
		index = {}
		for e in frappe.get_all("Employee", fields=["name", "employee_name"]):
			key = _norm_name(e.employee_name)
			if key:
				index.setdefault(key, []).append(e.name)
		for r in empty:
			hits = index.get(_norm_name(r.fio)) or []
			if len(hits) == 1:          # ikkita bir xil ism bo'lsa — qo'lda tanlansin
				r.employee = hits[0]


def _pick_sheet(wb, varaq=None):
	"""Ish haqi varag'ini topish: nomi berilsa o'sha, aks holda Ф.И.О.+Оклад bor varaq."""
	if varaq:
		if varaq not in wb.sheetnames:
			frappe.throw(_("Faylda '{0}' nomli varaq yo'q. Mavjudlari: {1}").format(
				varaq, ", ".join(wb.sheetnames)))
		return wb[varaq]

	for ws in wb.worksheets:
		for row in ws.iter_rows(min_row=1, max_row=12):
			heads = {_norm_header(c.value) for c in row if c.value}
			if any(h in ("ф.и.о.", "фио", "f.i.sh") for h in heads) and \
					any(h.startswith("оклад") or h == "oklad" for h in heads):
				return ws
	frappe.throw(_("Faylda ish haqi varag'i topilmadi (Ф.И.О. va Оклад ustunlari bo'lishi kerak)."))


def _header_row(ws):
	"""Sarlavha qatori va ustun raqamlari."""
	for r_idx, row in enumerate(ws.iter_rows(min_row=1, max_row=12), start=1):
		mapping = {}
		for i, c in enumerate(row):
			f = COLUMN_MAP.get(_norm_header(c.value))
			if f and f not in mapping:
				mapping[f] = i + 1
		if "fio" in mapping and "oklad" in mapping:
			return r_idx, mapping
	frappe.throw(_("Varaqda sarlavha qatori topilmadi."))


@frappe.whitelist()
def fill_from_file(docname):
	"""Biriktirilgan Excel fayldan vedomost qatorlarini to'ldirish.

	Har oy qayta bosish mumkin — eski qatorlar almashtiriladi. Qo'lda ulangan
	Employee bog'lamalari F.I.Sh bo'yicha saqlab qolinadi.
	"""
	doc = frappe.get_doc("Oylik Vedomost", docname)
	doc.check_permission("write")
	if not doc.fayl:
		frappe.throw(_("Avval Excel faylni biriktiring."))

	try:
		import openpyxl
	except ImportError:
		frappe.throw(_("openpyxl kutubxonasi o'rnatilmagan."))

	file_doc = frappe.get_doc("File", {"file_url": doc.fayl})
	wb = openpyxl.load_workbook(file_doc.get_full_path(), data_only=True, read_only=True)
	ws = _pick_sheet(wb, doc.varaq)
	hrow, cols = _header_row(ws)

	# oldingi qo'lda ulangan bog'lamalarni eslab qolamiz
	prev_links = {_norm_name(r.fio): r.employee for r in doc.qatorlar if r.employee}

	rows = []
	for row in ws.iter_rows(min_row=hrow + 1):
		# read_only rejimda bo'sh katak .column bermaydi — o'rin bo'yicha o'qiymiz
		cells = {i + 1: c.value for i, c in enumerate(row)}
		fio = cells.get(cols["fio"])
		if not fio or not str(fio).strip():
			continue
		item = {"fio": str(fio).strip()}
		for field, col in cols.items():
			if field == "fio":
				continue
			v = cells.get(col)
			if field in NUMERIC_FIELDS:
				item[field] = flt(v) if isinstance(v, (int, float)) else 0
			else:
				item[field] = str(v).strip() if v is not None else None
		# faqat sarlavha/jami kabi bo'sh qatorlarni tashlab yuboramiz
		if not any(flt(item.get(f)) for f in NUMERIC_FIELDS):
			continue
		item["employee"] = prev_links.get(_norm_name(item["fio"]))
		rows.append(item)

	wb.close()
	if not rows:
		frappe.throw(_("Varaqdan xodim qatorlari o'qilmadi — varaq nomini tekshiring."))

	doc.set("qatorlar", [])
	for r in rows:
		doc.append("qatorlar", r)
	doc.save()

	linked = sum(1 for r in doc.qatorlar if r.employee)
	return {
		"varaq": ws.title,
		"count": len(doc.qatorlar),
		"linked": linked,
		"unlinked": len(doc.qatorlar) - linked,
		"jami_oklad": doc.jami_oklad,
		"jami_nachisleniya": doc.jami_nachisleniya,
	}


def get_oklad_map(oy=None, to_date=None):
	"""Personal bo'limi uchun: xodim -> oklad (va ishlangan kun) xaritasi.

	`oy` berilmasa, `to_date` oyi olinadi; u ham topilmasa eng so'nggi vedomost.
	Kalitlar: Employee ID va normallashtirilgan F.I.Sh — ikkalasi ham qidiriladi.
	"""
	name = None
	if oy:
		name = frappe.db.get_value("Oylik Vedomost", {"oy": oy}, "name")
	if not name and to_date:
		name = frappe.db.get_value("Oylik Vedomost", {"oy": str(to_date)[:7]}, "name")
	if not name:
		rows = frappe.get_all("Oylik Vedomost", fields=["name"], order_by="oy desc", limit=1)
		name = rows[0].name if rows else None
	if not name:
		return {}, None

	out = {}
	for r in frappe.get_all("Oylik Vedomost Qatori", filters={"parent": name},
	                        fields=["fio", "employee", "oklad", "kun", "rejim", "doljnost"]):
		data = {"oklad": flt(r.oklad), "kun": flt(r.kun), "rejim": flt(r.rejim),
		        "doljnost": r.doljnost, "fio": r.fio}
		if r.employee:
			out[r.employee] = data
		key = _norm_name(r.fio)
		if key:
			out.setdefault("name::" + key, data)
	return out, name
