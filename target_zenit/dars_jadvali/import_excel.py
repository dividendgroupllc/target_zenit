# Copyright (c) 2026, Target Zenit
# Maktabning mavjud "Umumiy dars jadvali" Excel faylini tizimga ko'chirish.
#
# Fayl formati (barqaror):
#   A ustun = kun (birlashtirilgan katak), B = dars №, C = vaqt, D..W = 20 sinf
#   Yacheyka: "FAN\nO'qituvchi"  yoki  "FAN (N daraja)\nHigh: X\nMiddle: Y\nLow: Z"
#
# Bosqichlar (har biri idempotent):
#   1 fanlar (Course)  2 o'qituvchilar (Instructor)  3 guruh+blok  4 jadval yozuvlari
#
# Ishga tushirish:
#   bench --site <sayt> execute target_zenit.dars_jadvali.import_excel.import_all \
#       --kwargs "{'path':'/home/user/Downloads/Umumiy_dars_jadvali.xlsx','dry_run':1}"
from __future__ import annotations

import json
import re

import frappe

KUNLAR = ["Dushanba", "Seshanba", "Chorshanba", "Payshanba", "Juma", "Shanba"]

# Fan kodi -> to'liq nom (Excel "Izoh" varag'i va umumiy bilim asosida)
FAN_NOMLARI = {
	"MATH": "Matematika", "ENGLISH": "Ingliz tili", "RUS": "Rus tili",
	"UZBEK": "O'zbek tili", "CHINESE": "Xitoy tili", "SCIENCE": "Tabiiy fanlar",
	"PHYSICS": "Fizika", "GEO": "Geografiya", "HISTORY": "Tarix",
	"HIST.U": "O'zbekiston tarixi", "HIST.W": "Jahon tarixi", "IT": "Axborot texnologiyalari",
	"AI": "Sun'iy intellekt", "ROBOTICS": "Robototexnika", "CHESS": "Shaxmat",
	"ART": "Tasviriy san'at", "MUSIC": "Musiqa", "P.E / GYM": "Jismoniy tarbiya",
	"MENTAL.M": "Mental matematika", "MOR EDU": "Axloq tarbiyasi",
	"PER DEV": "Personal Development", "G PRES": "Group presentation",
	"FINANCE": "Moliya", "BUS": "Business", "ECO": "Economics",
	"HOMEWORK": "Uy vazifasi soati", "EXTRA": "Qo'shimcha dars", "CHOICE": "Tanlov darsi",
}

# Jadvaldagi qisqa ism -> Employee'dagi to'liq ism (transliteratsiya farqlari)
EMPLOYEE_MAP = {
	"Erdullaeva J.": "Yerdullayeva Janar",
	"Matyakubov L.": "Matyaqubov Lutfullo Habibullo o'g'li",
	"Karimxonov Abbos": "Karimkhonov Abbosxon Doniyorxon o'gli",
	"Abdug'aniyev F.": "Abduganiyev Firdavsbek Xabibullo o'gli",
	"Djiyenbayeva R.": "Djiyenbaeva Robiya Ilyasovna",
	"Mo'minjonov I.": "Mominjonov Ilyosbek Shuxrat o'gli",
	"Anvarovna X.": "Dadaboyeva Xosiyat Anvar qizi – 2 B",
	"Nigmatov M.": "Muhammadsodiq Nigmatov",
	"MR Dileep Mani": "Dileep Mani",
	"Rasulova M.": "Rasuleva Malika Olimovna",
	"Anorqulov Lazizbek": "Anorqul Lazizbek Alisher o'gli",
	"Botirova Surayyo Z.": "Botirova Surayyo",
	"Mamatov Diyor": "Mamatov Dior Orifovich",
	"Maulenova A.": "Maulenova Aydin Muxambetovna",
	"Muxammadjonova D.": "Durdona Muxammadjonova Rustamjon qizi",
	"Adilova N.": "Adilova Nigora Baxadirovna",
	"Adilova Nigora": "Adilova Nigora Baxadirovna",
	"Abdumajidova S.": "Abdumajidova Shoxsanam Ikromjon qizi",
	"Abdugulomova M.": "Abdugulomova Muxlisaxon No'monxon qizi",
	"Abdug'afforov A.": "Abdug'afforov Abdurashid Abduvali o'gli",
	"Fozilova N.": "Fozilova Nargiza Shuxrat qizi 1B",
	"Musulmonov M.": "Musulmonov Mamarajab Panji o'gli",
	"Usmonov I.": "Usmonov Ismoil",
	"Adxamov Nasrullo": "Adxamov Nasrullo Ikromjon o'g'li",
	"Boboyeva Sabrina": "Boboyeva Sabrina Lutfullayevna",
	"Saidazizova F.": "Saidazizova Farangiz Zinatilla qizi",
	"Kabulov J.": "Kabulov Jamshid Ilxomovich",
	"Lucille J. Ilao": "Lucille Johnalyn Ilao – 2 A",
	"Nida Umair": "Nida Umair – 4 A",
	"Glory Ilao": "Glory Ilao",
	"Poziljonova N.": "Polvonova Nigina",
}

# To'liq ismlar (skrinshotlardan) — Instructor nomi sifatida ishlatiladi
TOLIQ_ISM = {
	"Erdullaeva J.": "Erdullaeva Janar", "Matyakubov L.": "Matyakubov Lutfullo",
	"Abdumajidova S.": "Abdumajidova Shoxsanam", "Fozilova N.": "Fozilova Nargiza",
	"Abdugulomova M.": "Abdugulomova Muxlisaxon", "Musulmonov M.": "Musulmonov Mamarajab",
	"Anvarovna X.": "Anvarovna Xosiyat", "Adilova N.": "Adilova Nigora",
	"Abdug'aniyev F.": "Abdug'aniyev Firdavsbek", "Saidahmadov M.": "Saidahmadov Muhammadali",
	"Abdug'afforov A.": "Abdug'afforov Abdurashid", "Dushayev Fayoz": "Dushayev Fayoz Anvarovich",
	"Kabulov J.": "Kabulov Jamshid Ilhomovich", "Mamatov Diyor": "Mamatov Diyor Orifovich",
	"Usmonov I.": "Usmonov Ismoil", "Nigmatov M.": "Nigmatov Muhammadsodiq",
	"Poziljonova N.": "Poziljonova Nigina", "Saidazizova F.": "Saidazizova Farangiz",
	"Mo'minjonov I.": "Mo'minjonov Ilyosbek", "Djiyenbayeva R.": "Djiyenbayeva Robiya Ilyasovna",
	"Maulenova A.": "Maulenova Aydin Muhambetovna", "Saydaxmedova S.": "Saydaxmedova Sevara",
	"Muxammadjonova D.": "Muxammadjonova Durdona", "Lucille J. Ilao": "Lucille Johnalyn Ilao",
	"MR Dileep Mani": "Dileep Mani", "Rasulova M.": "Rasulova Malika",
	"Botirova Surayyo Z.": "Botirova Surayyo", "Muxtorov Husan": "Muxtorov Husan",
}

VAKANTLAR = {"AI VACANT", "Chess vakant", "Music vakant", "TUTOR", "Personal Development"}



def _fayl_yoli(path: str) -> str:
	"""Fayl yo'lini aniqlaydi. Uch xil ko'rinishni qabul qiladi:
	  - to'liq yo'l: /home/user/Downloads/Umumiy_dars_jadvali.xlsx
	  - Frappe File URL: /files/Umumiy_dars_jadvali.xlsx yoki /private/files/...
	  - faqat fayl nomi: Umumiy_dars_jadvali.xlsx  (File doctype'dan topiladi)
	Shu sababli serverga SSH bilan fayl ko'chirish shart emas — Desk'dan
	(/app/file) yuklab, nomini berish kifoya."""
	import os

	if path.startswith("/files/"):
		return frappe.get_site_path("public", "files", os.path.basename(path))
	if path.startswith("/private/files/"):
		return frappe.get_site_path("private", "files", os.path.basename(path))
	if os.path.isabs(path) and os.path.exists(path):
		return path

	# fayl nomi bo'yicha File doctype'dan qidirish (eng oxirgi yuklangani).
	# DIQQAT: Frappe bir xil nomli fayl bo'lsa nomiga hash qo'shadi
	# (Umumiy_dars_jadvali7735ee.xlsx) — shuning uchun LIKE bilan qidiramiz.
	nom = os.path.basename(path)
	asos, kengaytma = os.path.splitext(nom)
	f = frappe.db.get_value(
		"File", {"file_name": nom}, ["file_url", "is_private"], order_by="creation desc", as_dict=True
	)
	if not f:
		f = frappe.db.get_value(
			"File",
			{"file_name": ["like", f"{asos}%{kengaytma}"]},
			["file_url", "is_private"],
			order_by="creation desc",
			as_dict=True,
		)
	if not f:
		mavjud = frappe.get_all(
			"File",
			filters={"file_name": ["like", "%.xlsx"]},
			fields=["file_name", "file_url"],
			order_by="creation desc",
			limit_page_length=15,
		)
		royxat = "\n".join(f"  - {x.file_name}" for x in mavjud) or "  (hech qanday .xlsx yuklanmagan)"
		frappe.throw(
			f"Fayl topilmadi: {nom}\n\nServerda mavjud .xlsx fayllar:\n{royxat}\n\n"
			"Yechim: Desk'da /app/file sahifasiga faylni yuklang, "
			"yoki to'liq yo'lini bering (masalan /home/frappe/Umumiy_dars_jadvali.xlsx)."
		)
	papka = "private" if f.is_private else "public"
	return frappe.get_site_path(papka, "files", os.path.basename(f.file_url))


# ------------------------------------------------------------------ o'qish
def _oqish(path: str):
	import openpyxl

	wb = openpyxl.load_workbook(_fayl_yoli(path), data_only=True)
	ws = wb["Umumiy jadval"]
	sinflar = [ws.cell(4, c).value for c in range(4, ws.max_column + 1)]
	yacheykalar = []
	kun = None
	for r in range(5, 55):
		d = ws.cell(r, 1).value
		if d:
			kun = str(d).strip()
		no = ws.cell(r, 2).value
		for i, c in enumerate(range(4, ws.max_column + 1)):
			v = ws.cell(r, c).value
			if not v:
				continue
			satrlar = [p.strip() for p in str(v).split("\n") if p.strip()]
			yacheykalar.append(
				{
					"kun": kun,
					"dars": int(no),
					"sinf_qisqa": sinflar[i],
					"fan_matn": satrlar[0],
					"oqituvchilar": satrlar[1:],
				}
			)
	return sinflar, yacheykalar


def _sinf_tolik(qisqa: str) -> str | None:
	"""Excel '1A' -> bazadagi '1 A(full)'."""
	m = re.match(r"^\s*(\d+)\s*([AB])\s*$", str(qisqa))
	if not m:
		return None
	raqam, harf = m.group(1), m.group(2)
	for sg in frappe.get_all("Student Group", pluck="name"):
		mm = re.match(r"^\s*(\d+)\s*([AB])\b", sg)
		if mm and mm.group(1) == raqam and mm.group(2) == harf:
			return sg
	return None


def _fan_kodi(matn: str) -> str:
	return matn.split(" (")[0].strip()


def _daraja_soni(matn: str) -> int:
	m = re.search(r"\((\d+)\s*daraja\)", matn)
	return int(m.group(1)) if m else 0


# ------------------------------------------------------- 1. fanlar (Course)
def fanlar_yarat(yacheykalar, dry_run=0):
	kodlar = sorted({_fan_kodi(y["fan_matn"]) for y in yacheykalar})
	yangi = []
	for kod in kodlar:
		if frappe.db.exists("Course", kod):
			continue
		yangi.append(kod)
		if dry_run:
			continue
		frappe.get_doc(
			{
				"doctype": "Course",
				"course_name": kod,
				"course_code": kod,
				"description": FAN_NOMLARI.get(kod, ""),
			}
		).insert(ignore_permissions=True)
	return {"jami": len(kodlar), "yangi": yangi}


# ------------------------------------------- 2. o'qituvchilar (Instructor)
def _oqituvchi_nomi(xom: str) -> str:
	"""'High: Matyakubov L.' -> 'Matyakubov L.'"""
	return xom.split(":", 1)[-1].strip() if ":" in xom else xom.strip()


def oqituvchilar_yarat(yacheykalar, dry_run=0):
	nomlar = sorted({_oqituvchi_nomi(t) for y in yacheykalar for t in y["oqituvchilar"]})
	natija = {"jami": len(nomlar), "yangi": [], "employee_bilan": [], "employee_siz": [], "vakant": []}
	for nom in nomlar:
		toliq = TOLIQ_ISM.get(nom, nom)
		vakant = nom in VAKANTLAR
		emp_nomi = EMPLOYEE_MAP.get(nom)
		employee = None
		if emp_nomi:
			employee = frappe.db.get_value("Employee", {"employee_name": emp_nomi}, "name")
		if not employee and not vakant:
			employee = frappe.db.get_value("Employee", {"employee_name": toliq}, "name")

		mavjud = frappe.db.get_value("Instructor", {"instructor_name": toliq}, "name")
		if vakant:
			natija["vakant"].append(toliq)
		elif employee:
			natija["employee_bilan"].append(f"{toliq} -> {employee}")
		else:
			natija["employee_siz"].append(toliq)

		if mavjud or dry_run:
			continue
		doc = frappe.get_doc(
			{
				"doctype": "Instructor",
				"instructor_name": toliq,
				"employee": employee,
				"custom_vakant": 1 if vakant else 0,
			}
		)
		doc.flags.ignore_mandatory = True
		doc.insert(ignore_permissions=True)
		natija["yangi"].append(doc.name)
	return natija


def _instructor(nom_xom: str) -> str | None:
	nom = _oqituvchi_nomi(nom_xom)
	toliq = TOLIQ_ISM.get(nom, nom)
	return frappe.db.get_value("Instructor", {"instructor_name": toliq}, "name")


# ------------------------------------------- 3. guruhlar va bloklar
def _slot_bloklari(yacheykalar):
	"""(kun, dars, fan_matn) -> blok kaliti (fan_kodi, sinflar tuple).

	MUHIM: bitta sinf bir fan bo'yicha turli vaqtda turli blokda bo'lishi mumkin
	(masalan ENGLISH: 3A yolg'iz, yoki 3A+5A+5B birga) — shuning uchun blok
	har bir vaqt sloti uchun alohida aniqlanadi."""
	from collections import defaultdict

	vaqt_boyicha = defaultdict(set)
	for y in yacheykalar:
		if _daraja_soni(y["fan_matn"]):
			vaqt_boyicha[(y["kun"], y["dars"], y["fan_matn"])].add(y["sinf_qisqa"])
	slot_kalit = {}
	for slot, sinflar in vaqt_boyicha.items():
		kod = _fan_kodi(slot[2])
		slot_kalit[slot] = (kod, tuple(sorted(sinflar, key=lambda s: (len(s), s))))
	return slot_kalit


def _qoshma_slotlar(yacheykalar):
	"""Qo'shma darslar: bir o'qituvchi + bir fan + bir vaqtda bir nechta sinf
	(masalan jismoniy tarbiya 1A+1B birga). Bu konflikt emas — bitta dars.

	Qaytaradi: (kun, dars, fan_matn, o'qituvchi) -> sinflar tuple."""
	from collections import defaultdict

	gr = defaultdict(list)
	for y in yacheykalar:
		if _daraja_soni(y["fan_matn"]):
			continue
		for t in y["oqituvchilar"]:
			gr[(y["kun"], y["dars"], y["fan_matn"], t)].append(y["sinf_qisqa"])
	return {
		k: tuple(sorted(v, key=lambda s: (len(s), s)))
		for k, v in gr.items()
		if len(v) > 1
	}


def _qoshma_bloklar(yacheykalar):
	"""Qo'shma dars bloklari: (fan_kodi, sinflar, o'qituvchi) -> necha marta."""
	from collections import Counter

	c = Counter()
	for (_, _, fan_matn, oqit), sinflar in _qoshma_slotlar(yacheykalar).items():
		c[(_fan_kodi(fan_matn), sinflar, _oqituvchi_nomi(oqit))] += 1
	return c


def _blok_kaliti(yacheykalar):
	"""Barqaror bloklar ro'yxati: kalit -> haftada necha marta uchraydi."""
	from collections import Counter

	return Counter(_slot_bloklari(yacheykalar).values())


def _blok_nomi(kod, sinflar):
	raqamlar = sorted({re.match(r"\d+", s).group(0) for s in sinflar if re.match(r"\d+", s)}, key=int)
	oraliq = raqamlar[0] if len(raqamlar) == 1 else f"{raqamlar[0]}-{raqamlar[-1]}"
	return f"{kod} {oraliq}"


def guruhlar_yarat(yacheykalar, academic_year, dry_run=0):
	"""Daraja/tanlov guruhlari (Student Group) va bloklar (Dars Bloki)."""
	from collections import defaultdict

	bloklar = _blok_kaliti(yacheykalar)
	slot_kalit = _slot_bloklari(yacheykalar)
	# blok -> {daraja: o'qituvchi} — slot kaliti orqali aniq bog'lanadi
	darajalar = defaultdict(dict)
	for y in yacheykalar:
		if not _daraja_soni(y["fan_matn"]):
			continue
		kalit = slot_kalit.get((y["kun"], y["dars"], y["fan_matn"]))
		if not kalit:
			continue
		for t in y["oqituvchilar"]:
			daraja = t.split(":", 1)[0].strip() if ":" in t else "—"
			darajalar[kalit][daraja] = _oqituvchi_nomi(t)

	natija = {"bloklar": [], "guruhlar": [], "sinf_topilmadi": set()}
	for kalit, marta in sorted(bloklar.items(), key=lambda x: -x[1]):
		kod, sinflar = kalit
		nomi = _blok_nomi(kod, sinflar)
		# bir xil nomli ikki blok bo'lsa (masalan IT 5-6 va IT 5-6) — farqlash
		if nomi in [b["nomi"] for b in natija["bloklar"]]:
			nomi = f"{nomi} ({'+'.join(sinflar)})"
		toliq_sinflar = []
		for s in sinflar:
			t = _sinf_tolik(s)
			if t:
				toliq_sinflar.append(t)
			else:
				natija["sinf_topilmadi"].add(s)

		guruh_qatorlari = []
		for daraja, oqituvchi in sorted(darajalar[kalit].items()):
			guruh_nomi = f"{kod} {daraja} {_blok_nomi(kod, sinflar).split(' ', 1)[-1]}"
			guruh_qatorlari.append(
				{"daraja": daraja, "guruh_nomi": guruh_nomi, "oqituvchi": oqituvchi}
			)
			natija["guruhlar"].append(guruh_nomi)
		natija["bloklar"].append(
			{"nomi": nomi, "fan": kod, "sinflar": toliq_sinflar, "guruhlar": guruh_qatorlari, "haftada": marta}
		)

	# --- qo'shma darslar (A+B birga): bitta guruhli blok ---
	for (kod, sinflar, oqituvchi), marta in _qoshma_bloklar(yacheykalar).items():
		toliq = [_sinf_tolik(s) for s in sinflar]
		if not all(toliq):
			natija["sinf_topilmadi"].update(s for s in sinflar if not _sinf_tolik(s))
			continue
		tag = "+".join(sinflar)
		nomi = f"{kod.replace('/', '-').strip()} {tag} (qo'shma)"
		natija["bloklar"].append(
			{
				"nomi": nomi,
				"fan": kod,
				"sinflar": toliq,
				"qoshma": 1,
				"guruhlar": [{"daraja": "Qo'shma", "guruh_nomi": nomi, "oqituvchi": oqituvchi}],
				"haftada": marta,
			}
		)
		natija["guruhlar"].append(nomi)

	if dry_run:
		natija["sinf_topilmadi"] = sorted(natija["sinf_topilmadi"])
		return natija

	# --- yozish ---
	for b in natija["bloklar"]:
		for g in b["guruhlar"]:
			if not frappe.db.exists("Student Group", g["guruh_nomi"]):
				sg = frappe.get_doc(
					{
						"doctype": "Student Group",
						"student_group_name": g["guruh_nomi"],
						"group_based_on": "Activity" if b["fan"] in ("CHOICE", "EXTRA") else "Course",
						"course": None if b["fan"] in ("CHOICE", "EXTRA") else b["fan"],
						"academic_year": academic_year,
						"custom_guruh_turi": (
							"Qo'shma" if b.get("qoshma")
							else "Tanlov" if b["fan"] in ("CHOICE", "EXTRA") else "Daraja"
						),
						"custom_daraja": g["daraja"],
						"custom_azo_sinflar": ", ".join(b["sinflar"]),
					}
				)
				sg.flags.ignore_mandatory = True
				sg.insert(ignore_permissions=True)
		if not frappe.db.exists("Dars Bloki", b["nomi"]):
			frappe.get_doc(
				{
					"doctype": "Dars Bloki",
					"nomi": b["nomi"],
					"fan": b["fan"] if frappe.db.exists("Course", b["fan"]) else None,
					"academic_year": academic_year,
					"sinflar": [{"sinf": s} for s in b["sinflar"]],
					"guruhlar": [
						{
							"daraja": g["daraja"],
							"guruh": g["guruh_nomi"],
							"oqituvchi": _instructor(g["oqituvchi"]),
						}
						for g in b["guruhlar"]
					],
				}
			).insert(ignore_permissions=True)
		for g in b["guruhlar"]:
			frappe.db.set_value("Student Group", g["guruh_nomi"], "custom_blok", b["nomi"], update_modified=False)

	natija["sinf_topilmadi"] = sorted(natija["sinf_topilmadi"])
	return natija


# ------------------------------------------- 4. jadval yozuvlari
def yozuvlar_yarat(yacheykalar, versiya, dry_run=0, xotira_bloklar=None):
	"""Yacheykalardan Jadval Yozuvi. Daraja darslari blok guruhlariga yoziladi
	(sinf ustunlari bo'yicha takrorlanmaydi), oddiy darslar — sinf guruhiga."""
	from collections import defaultdict

	natija = {"yaratildi": 0, "bor_edi": 0, "xato": [], "otkazildi": []}
	korilgan_bloklar = set()
	slot_kalit = _slot_bloklari(yacheykalar)
	qoshma = _qoshma_slotlar(yacheykalar)

	# blok kaliti (fan, sinflar) -> bazadagi Dars Bloki nomi
	blok_nomi_cache = {}

	def blok_top(kalit):
		if kalit in blok_nomi_cache:
			return blok_nomi_cache[kalit]
		kod, sinflar = kalit
		if xotira_bloklar is not None:
			# dry_run: bloklar hali bazada yo'q — xotiradagi ro'yxatdan topamiz
			toliq_set = {_sinf_tolik(x) for x in sinflar}
			for b in xotira_bloklar:
				if b["fan"] == kod and set(b["sinflar"]) == toliq_set:
					blok_nomi_cache[kalit] = b["nomi"]
					return b["nomi"]
			blok_nomi_cache[kalit] = None
			return None
		toliq = {_sinf_tolik(s) for s in sinflar}
		topildi = None
		for b in frappe.get_all("Dars Bloki", filters={"fan": kod}, pluck="name"):
			b_sinflar = set(
				frappe.get_all("Dars Bloki Sinfi", filters={"parent": b}, pluck="sinf")
			)
			if b_sinflar == toliq:
				topildi = b
				break
		blok_nomi_cache[kalit] = topildi
		return topildi

	for y in sorted(yacheykalar, key=lambda x: (KUNLAR.index(x["kun"]), x["dars"])):
		fan_kod = _fan_kodi(y["fan_matn"])
		sinf = _sinf_tolik(y["sinf_qisqa"])
		if not sinf:
			natija["otkazildi"].append(f"{y['sinf_qisqa']} sinfi topilmadi")
			continue

		if _daraja_soni(y["fan_matn"]):
			# Daraja darsi: blok bo'yicha bir marta (slot kaliti orqali aniq blok)
			kalit = slot_kalit.get((y["kun"], y["dars"], y["fan_matn"]))
			blok = blok_top(kalit) if kalit else None
			if not blok:
				natija["xato"].append(f"{y['kun']} {y['dars']} {y['sinf_qisqa']}: blok topilmadi ({fan_kod})")
				continue
			slot = (y["kun"], y["dars"], blok)
			if slot in korilgan_bloklar:
				continue
			korilgan_bloklar.add(slot)
			if xotira_bloklar is not None:
				xb = next((b for b in xotira_bloklar if b["nomi"] == blok), None)
				guruhlar = [
					frappe._dict(guruh=g["guruh_nomi"], oqituvchi=g["oqituvchi"])
					for g in (xb["guruhlar"] if xb else [])
				]
			else:
				guruhlar = frappe.get_all(
					"Dars Bloki Guruhi", filters={"parent": blok},
					fields=["guruh", "oqituvchi"], order_by="idx",
				)
			for g in guruhlar:
				_yoz(versiya, y, g.guruh, fan_kod, g.oqituvchi, blok, natija, dry_run)
		else:
			oqituvchi_xom = y["oqituvchilar"][0] if y["oqituvchilar"] else None
			oqituvchi = _instructor(oqituvchi_xom) if oqituvchi_xom else None

			# Qo'shma dars (A+B birga) — bitta guruhga bitta yozuv
			q_sinflar = qoshma.get((y["kun"], y["dars"], y["fan_matn"], oqituvchi_xom)) if oqituvchi_xom else None
			if q_sinflar:
				q_nomi = f"{fan_kod.replace('/', '-').strip()} {'+'.join(q_sinflar)} (qo'shma)"
				slot = (y["kun"], y["dars"], q_nomi)
				if slot in korilgan_bloklar:
					continue
				korilgan_bloklar.add(slot)
				q_blok = q_nomi if (dry_run or frappe.db.exists("Dars Bloki", q_nomi)) else None
				_yoz(versiya, y, q_nomi, fan_kod, oqituvchi, q_blok, natija, dry_run)
				continue

			_yoz(versiya, y, sinf, fan_kod, oqituvchi, None, natija, dry_run)

	return natija


def _yoz(versiya, y, guruh, fan_kod, oqituvchi, blok, natija, dry_run):
	mavjud = frappe.db.exists(
		"Jadval Yozuvi",
		{"versiya": versiya, "kun": y["kun"], "dars_raqami": y["dars"], "guruh": guruh},
	)
	if mavjud:
		natija["bor_edi"] += 1
		return
	if dry_run:
		natija["yaratildi"] += 1
		return
	try:
		frappe.get_doc(
			{
				"doctype": "Jadval Yozuvi",
				"versiya": versiya,
				"kun": y["kun"],
				"dars_raqami": y["dars"],
				"guruh": guruh,
				"fan": fan_kod if frappe.db.exists("Course", fan_kod) else None,
				"oqituvchi": oqituvchi,
				"blok": blok,
			}
		).insert(ignore_permissions=True)
		natija["yaratildi"] += 1
	except Exception as e:
		natija["xato"].append(f"{y['kun']} {y['dars']} {guruh}: {str(e)[:150]}")


# ------------------------------------------------------------- boshqaruv
def import_all(
	path="/home/user/Downloads/Umumiy_dars_jadvali.xlsx",
	versiya_nomi="2026-2027 asosiy",
	academic_year=None,
	dry_run=1,
):
	"""To'liq import. dry_run=1 — hech narsa yozilmaydi, faqat hisobot."""
	academic_year = academic_year or frappe.db.get_value("Academic Year", {}, "name")
	sinflar, yacheykalar = _oqish(path)
	hisobot = {
		"fayl": path,
		"yacheykalar": len(yacheykalar),
		"sinflar": len(sinflar),
		"dry_run": bool(dry_run),
	}

	hisobot["1_fanlar"] = fanlar_yarat(yacheykalar, dry_run)
	hisobot["2_oqituvchilar"] = oqituvchilar_yarat(yacheykalar, dry_run)
	hisobot["3_guruhlar"] = guruhlar_yarat(yacheykalar, academic_year, dry_run)

	# versiya
	if not dry_run and not frappe.db.exists("Jadval Versiyasi", versiya_nomi):
		frappe.get_doc(
			{
				"doctype": "Jadval Versiyasi",
				"nomi": versiya_nomi,
				"academic_year": academic_year,
				"holat": "Qoralama",
				"amal_boshlanishi": f"{str(academic_year).split('-')[0]}-09-01",
				"izoh": f"Excel'dan import: {path}",
			}
		).insert(ignore_permissions=True)

	hisobot["4_yozuvlar"] = yozuvlar_yarat(
		yacheykalar, versiya_nomi, dry_run,
		xotira_bloklar=hisobot["3_guruhlar"]["bloklar"] if dry_run else None,
	)

	if not dry_run:
		frappe.db.commit()

	qisqa = {
		"yacheykalar": hisobot["yacheykalar"],
		"fanlar_yangi": len(hisobot["1_fanlar"]["yangi"]),
		"oqituvchilar": {
			"jami": hisobot["2_oqituvchilar"]["jami"],
			"employee_bilan": len(hisobot["2_oqituvchilar"]["employee_bilan"]),
			"employee_siz": hisobot["2_oqituvchilar"]["employee_siz"],
			"vakant": hisobot["2_oqituvchilar"]["vakant"],
		},
		"bloklar": len(hisobot["3_guruhlar"]["bloklar"]),
		"daraja_guruhlari": len(set(hisobot["3_guruhlar"]["guruhlar"])),
		"sinf_topilmadi": hisobot["3_guruhlar"]["sinf_topilmadi"],
		"yozuvlar": {
			"yaratildi": hisobot["4_yozuvlar"]["yaratildi"],
			"bor_edi": hisobot["4_yozuvlar"]["bor_edi"],
			"xato_soni": len(hisobot["4_yozuvlar"]["xato"]),
			"xato_namuna": hisobot["4_yozuvlar"]["xato"][:10],
		},
	}
	print(json.dumps(qisqa, ensure_ascii=False, indent=1))
	return qisqa


def bloklar_royxati(path="/home/user/Downloads/Umumiy_dars_jadvali.xlsx"):
	"""Faqat ko'rish: qanday bloklar aniqlanadi."""
	_, yacheykalar = _oqish(path)
	r = guruhlar_yarat(yacheykalar, None, dry_run=1)
	for b in r["bloklar"]:
		print(f"{b['haftada']:3}x  {b['nomi']:22} <- {', '.join(b['sinflar'])}")
		for g in b["guruhlar"]:
			print(f"          {g['daraja']:14} {g['guruh_nomi']:26} {g['oqituvchi']}")
	return r


def tozalash(versiya="2026-2027 asosiy"):
	"""Shu versiyaning barcha jadval yozuvlarini o'chiradi (qayta import uchun)."""
	n = 0
	for y in frappe.get_all("Jadval Yozuvi", filters={"versiya": versiya}, pluck="name"):
		frappe.delete_doc("Jadval Yozuvi", y, force=1, ignore_permissions=True)
		n += 1
	frappe.db.commit()
	print(f"O'chirildi: {n} ta yozuv ({versiya})")


def fayllar():
	"""Serverda qanday .xlsx fayllar yuklanganini ko'rsatadi (import oldidan tekshirish)."""
	rows = frappe.get_all(
		"File",
		filters={"file_name": ["like", "%.xlsx"]},
		fields=["file_name", "file_url", "is_private", "creation", "attached_to_doctype"],
		order_by="creation desc",
		limit_page_length=30,
	)
	if not rows:
		print("Hech qanday .xlsx fayl yuklanmagan. /app/file sahifasidan yuklang.")
		return rows
	print(f"{'FAYL NOMI':46} {'MAXFIY':7} YUKLANGAN")
	for r in rows:
		print(f"{r.file_name[:45]:46} {'ha' if r.is_private else 'yo`q':7} {str(r.creation)[:16]}")
	return rows


# ------------------------------------------------------------- UI import
IMPORT_ROLLAR = {"system manager", "zavuch", "academics user", "education manager"}


def _import_guard():
	if not (IMPORT_ROLLAR & {r.lower() for r in frappe.get_roles()}):
		frappe.throw("Jadval import qilish uchun ruxsat yo'q.", frappe.PermissionError)


@frappe.whitelist()
def ui_import(file_url: str, versiya_nomi: str = "2026-2027 asosiy", dry_run=1, tozalash_avval=0):
	"""Desk'dan import: faylni fon vazifasida o'qiydi (1-2 daqiqa).
	Natija realtime orqali qaytadi + keshda saqlanadi."""
	_import_guard()
	if not file_url:
		frappe.throw("Avval Excel faylni yuklang.")
	frappe.enqueue(
		"target_zenit.dars_jadvali.import_excel._ui_import_job",
		queue="long",
		timeout=1800,
		user=frappe.session.user,
		file_url=file_url,
		versiya_nomi=versiya_nomi,
		dry_run=int(dry_run or 0),
		tozalash_avval=int(tozalash_avval or 0),
	)
	return {"ok": 1}


def _ui_import_job(user, file_url, versiya_nomi, dry_run, tozalash_avval):
	natija = {}
	try:
		if tozalash_avval and not dry_run:
			tozalash(versiya_nomi)
		natija = import_all(path=file_url, versiya_nomi=versiya_nomi, dry_run=dry_run)
		natija["holat"] = "tugadi"
	except Exception as e:
		frappe.log_error(frappe.get_traceback(), "jadval import (UI)")
		natija = {"holat": "xato", "xabar": str(e)[:500]}

	natija["dry_run"] = int(dry_run)
	frappe.cache().set_value(f"jadval_import_natija:{user}", natija, expires_in_sec=3600)
	frappe.publish_realtime("jadval_import_tugadi", natija, user=user)


@frappe.whitelist()
def oxirgi_natija():
	"""Oxirgi import natijasi (sahifa yangilansa ham ko'rinsin)."""
	_import_guard()
	return frappe.cache().get_value(f"jadval_import_natija:{frappe.session.user}")
